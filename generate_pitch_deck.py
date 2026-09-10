import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation (16:9 Widescreen)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Exact Theme of Prime HRIS Web App)
BG_DARK = RGBColor(11, 15, 25)         # #0B0F19 (Deep Obsidian Slate)
CARD_BG = RGBColor(19, 27, 46)         # #131B2E (Dark Slate Card)
CARD_BG_ALT = RGBColor(26, 36, 61)     # #1A243D (Lighter Dark Card)
BORDER_DARK = RGBColor(38, 51, 80)     # #263350 (Subtle Slate Border)
ORANGE = RGBColor(255, 107, 0)         # #FF6B00 (Primary Industrial Orange)
ORANGE_LIGHT = RGBColor(255, 133, 51)  # #FF8533 (Light Orange)
ORANGE_PALE = RGBColor(255, 243, 235)  # #FFF3EB
MINT = RGBColor(0, 226, 176)          # #00E2B0 (Verified Emerald / Mint)
MINT_DARK = RGBColor(6, 78, 59)        # #064E3B
WHITE = RGBColor(255, 255, 255)
SLATE_MUTED = RGBColor(148, 163, 184)  # #94A3B8
SLATE_LIGHT = RGBColor(226, 232, 240)  # #E2E8F0
RED_ACCENT = RGBColor(244, 63, 94)     # #F43F5E (Alert Rose)

FONT_TITLE = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_CODE = "Consolas"

blank_slide_layout = prs.slide_layouts[6]

def apply_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.color.rgb = BG_DARK
    
    # Top Accent Strip (Web-style gradient bar simulation)
    top_strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    top_strip.fill.solid()
    top_strip.fill.fore_color.rgb = ORANGE
    top_strip.line.color.rgb = ORANGE

def add_header(slide, badge_text, title_text, subtitle_text=""):
    # Badge Pill
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.4), Inches(0.36))
    badge.fill.solid()
    badge.fill.fore_color.rgb = CARD_BG
    badge.line.color.rgb = ORANGE
    badge.line.width = Pt(1.5)
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.LEFT
    run_b = p_b.add_run()
    run_b.text = "  [ " + badge_text.upper() + " ]"
    run_b.font.name = FONT_CODE
    run_b.font.size = Pt(9.5)
    run_b.font.bold = True
    run_b.font.color.rgb = ORANGE_LIGHT
    
    # Title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(11.7), Inches(0.75))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    if subtitle_text:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = SLATE_MUTED

def add_footer(slide, current_slide, total_slides=12):
    # Bottom separator line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER_DARK
    line.line.color.rgb = BORDER_DARK
    
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"Prime HRIS Enterprise • Proposal Penawaran Sistem PT Dwi Martha Jaya | Hal {current_slide} dari {total_slides}"
    p.font.name = FONT_CODE
    p.font.size = Pt(8.5)
    p.font.color.rgb = SLATE_MUTED

def add_card(slide, x, y, w, h, border_color=BORDER_DARK, bg_color=CARD_BG):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    return card

# ==============================================================================
# SLIDE 1: COVER SLIDE
# ==============================================================================
s1 = prs.slides.add_slide(blank_slide_layout)
apply_background(s1)

card_cov = add_card(s1, 0.8, 0.8, 11.733, 5.8, border_color=ORANGE, bg_color=CARD_BG)

badge_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(1.3), Inches(5.2), Inches(0.42))
badge_box.fill.solid()
badge_box.fill.fore_color.rgb = RGBColor(30, 41, 65)
badge_box.line.color.rgb = ORANGE
badge_box.line.width = Pt(1.5)
tf = badge_box.text_frame
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = " [ SAMPLE PROJECT & DOKUMEN PENAWARAN RESMI ]"
r.font.name = FONT_CODE
r.font.size = Pt(10)
r.font.bold = True
r.font.color.rgb = ORANGE_LIGHT

tb = s1.shapes.add_textbox(Inches(1.3), Inches(1.9), Inches(10.5), Inches(2.2))
tf = tb.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = 0

p = tf.paragraphs[0]
p.text = "PRIME HRIS ENTERPRISE"
p.font.name = FONT_TITLE
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = WHITE

p2 = tf.add_paragraph()
p2.text = "Sistem Manajemen SDM Terpadu, Presensi Kamera Watermark Digital,\nSequential WiFi Gateway, & Otomasi Penggajian Berbasis Proyek Industri"
p2.font.name = FONT_BODY
p2.font.size = Pt(15)
p2.font.color.rgb = ORANGE_LIGHT

p3 = tf.add_paragraph()
p3.space_before = Pt(8)
p3.text = "Studi Solusi, Proof of Concept & Penawaran Komersial Khusus untuk Direksi & Manajemen PT Dwi Martha Jaya"
p3.font.name = FONT_BODY
p3.font.size = Pt(12)
p3.font.color.rgb = SLATE_LIGHT

meta_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(4.5), Inches(10.7), Inches(1.7))
meta_card.fill.solid()
meta_card.fill.fore_color.rgb = RGBColor(14, 20, 36)
meta_card.line.color.rgb = BORDER_DARK

tb_meta = s1.shapes.add_textbox(Inches(1.5), Inches(4.6), Inches(10.3), Inches(1.5))
tf_m = tb_meta.text_frame
tf_m.word_wrap = True

p_m1 = tf_m.paragraphs[0]
p_m1.text = "KLIEN SASARAN: PT DWI MARTHA JAYA | VENDOR PENGEMBANG: PT PRIME INFINITY SYSTEMS (PRIME PROJECTX)"
p_m1.font.name = FONT_CODE
p_m1.font.size = Pt(10)
p_m1.font.bold = True
p_m1.font.color.rgb = MINT

p_m2 = tf_m.add_paragraph()
p_m2.text = "Author & Platform Architect : Galih Primananda, S.E. (Email: galih@primeprojectx.net • WA: +62 895 2425 7778)\n" \
            "Nomor Dokumen               : PIS/PROP/DMJ-HRIS/2026/09 • Tanggal: 8 September 2026\n" \
            "Platform Showcase Live      : https://dmj.primeprojectx.net • https://primeprojectx.net"
p_m2.font.name = FONT_CODE
p_m2.font.size = Pt(9.5)
p_m2.font.color.rgb = SLATE_MUTED

# ==============================================================================
# SLIDE 2: EXECUTIVE SUMMARY & PROBLEM STATEMENT
# ==============================================================================
s2 = prs.slides.add_slide(blank_slide_layout)
apply_background(s2)
add_header(s2, "Executive Summary", "Latar Belakang & Analisis Masalah Operasional Proyek",
           "Tantangan kedisiplinan dan akuntabilitas jam kerja industri workshop & site konstruksi PT Dwi Martha Jaya")
add_footer(s2, 2)

problems = [
    ("Titip Absen & Fake GPS",
     "Presensi konvensional atau finger print statis di kantor tidak dapat memverifikasi kehadiran teknisi lapangan secara visual dan rentan manipulasi aplikasi lokasi fiktif.",
     "Risiko: Pembayaran gaji & lembur tidak valid", RED_ACCENT),
    ("Kekakuan Saat Tugas Darurat",
     "Teknisi yang ditugaskan mendadak ke site klien (Dinas Luar) terganjal radius geofence kantor sehingga sering dicatat alpa atau telat oleh sistem kaku.",
     "Risiko: Hambatan operasional & demotivasi staf", ORANGE),
    ("Payroll Gelondongan (No Costing)",
     "Gaji teknisi dibebankan secara total ke overhead tanpa dialokasikan ke nomor kontrak proyek spesifik (seperti Smelter Manyar JIIPE atau Fabrikasi Workshop).",
     "Risiko: Laporan laba/rugi proyek bias", ORANGE),
    ("Pembengkakan Biaya Lembur",
     "Tidak adanya formula baku Depnaker (1/173) yang otomatis serta absennya deteksi dini pembengkakan upah lembur (Labor Cost Overrun) sebelum proyek selesai.",
     "Risiko: Kerugian finansial tak terdeteksi", RED_ACCENT)
]

for i, (title, desc, impact, col) in enumerate(problems):
    row = i // 2
    col_idx = i % 2
    x = 0.8 + col_idx * 5.95
    y = 1.65 + row * 2.5
    
    add_card(s2, x, y, 5.75, 2.3, border_color=BORDER_DARK)
    
    tb = s2.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.18), Inches(5.25), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = f"MASALAH {i+1} : {title.upper()}"
    p.font.name = FONT_CODE
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = col
    
    p_desc = tf.add_paragraph()
    p_desc.space_before = Pt(4)
    p_desc.text = desc
    p_desc.font.name = FONT_BODY
    p_desc.font.size = Pt(10.5)
    p_desc.font.color.rgb = SLATE_LIGHT
    
    p_imp = tf.add_paragraph()
    p_imp.space_before = Pt(6)
    p_imp.text = f"Dampak: {impact}"
    p_imp.font.name = FONT_CODE
    p_imp.font.size = Pt(9)
    p_imp.font.bold = True
    p_imp.font.color.rgb = col

sol_card = add_card(s2, 0.8, 6.2, 11.733, 0.65, border_color=MINT, bg_color=CARD_BG_ALT)
tb_sol = s2.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_s = tb_sol.text_frame
p_s = tf_s.paragraphs[0]
r_s = p_s.add_run()
r_s.text = "[ SOLUSI PRIME HRIS ] : "
r_s.font.bold = True
r_s.font.color.rgb = MINT
r_s.font.name = FONT_CODE
r_s.font.size = Pt(10.5)
r_s2 = p_s.add_run()
r_s2.text = "Mengintegrasikan Presensi Kamera Watermark, Sequential WiFi Gate, Alokasi Biaya Proyek & Slip Depnaker dalam 1 platform web 100% otomatis."
r_s2.font.color.rgb = WHITE
r_s2.font.name = FONT_BODY
r_s2.font.size = Pt(10.5)

# ==============================================================================
# SLIDE 3: SYSTEM ARCHITECTURE & 5 CORE PILLARS
# ==============================================================================
s3 = prs.slides.add_slide(blank_slide_layout)
apply_background(s3)
add_header(s3, "Arsitektur Sistem", "5 Pilar Utama Ekosistem Prime HRIS Enterprise",
           "Alur otomatisasi terintegrasi dari presensi biometrik harian hingga laporan audit penggajian resmi")
add_footer(s3, 3)

pillars = [
    ("PILAR 1", "Presensi Kamera & Watermark", "Selfie webcam dengan stempel permanen NIK, Nama, Waktu WIB, GPS satelit, dan IP subnet. 100% anti titip absen."),
    ("PILAR 2", "Sequential WiFi Gateway", "Verifikasi berjenjang SSID DMJ-Corporate-5G & MAC Address hardware HP/Laptop. Blokir otomatis VPN & Fake GPS."),
    ("PILAR 3", "Database SDM & Kompensasi", "Sentralisasi NIK resmi, rekening bank payroll BCA/Mandiri, struktur gaji pokok, tunjangan, dan saldo cuti tahunan."),
    ("PILAR 4", "Payroll Depnaker & EAC", "Perhitungan upah lembur 1/173 baku Menakertrans, BPJS, PPh21, dan alokasi biaya tenaga kerja langsung ke proyek."),
    ("PILAR 5", "Multi-Tier Approval Hub", "Satu pintu persetujuan cuti, dispensasi keterlambatan, lembur, dinas luar mendadak, dan klaim reimbursement.")
]

for i, (tag, title, desc) in enumerate(pillars):
    x = 0.8 + i * 2.38
    card = add_card(s3, x, 1.7, 2.25, 4.3, border_color=ORANGE if i==0 or i==3 else BORDER_DARK)
    
    tb = s3.shapes.add_textbox(Inches(x + 0.15), Inches(1.85), Inches(1.95), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = tag
    p.font.name = FONT_CODE
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    p2.text = title
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    
    p3 = tf.add_paragraph()
    p3.space_before = Pt(12)
    p3.text = desc
    p3.font.name = FONT_BODY
    p3.font.size = Pt(10)
    p3.font.color.rgb = SLATE_LIGHT

stat_card = add_card(s3, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK, bg_color=CARD_BG_ALT)
tb_st = s3.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_st = tb_st.text_frame
p_st = tf_st.paragraphs[0]
p_st.text = "[ VALUE PROPOSITION ] : Memangkas 90% waktu rekap absensi bulanan | 0% kebocoran jam lembur fiktif | 100% kepatuhan audit Depnaker & Pajak."
p_st.font.name = FONT_CODE
p_st.font.size = Pt(10)
p_st.font.bold = True
p_st.font.color.rgb = MINT

# ==============================================================================
# SLIDE 4: MODUL PRESENSI KAMERA & WATERMARK DIGITAL
# ==============================================================================
s4 = prs.slides.add_slide(blank_slide_layout)
apply_background(s4)
add_header(s4, "Modul 1 • Presensi Biometrik", "Presensi Kamera Selfie & Watermark Digital Permanen",
           "Autentikasi kehadiran visual dengan stempel metadata kriptografis tak terhapuskan")
add_footer(s4, 4)

add_card(s4, 0.8, 1.7, 6.8, 4.3, border_color=BORDER_DARK)
tb_l = s4.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(6.3), Inches(3.9))
tf_l = tb_l.text_frame
tf_l.word_wrap = True

p = tf_l.paragraphs[0]
p.text = "FITUR UNGGULAN PRESENSI KAMERA CERDAS"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ORANGE

bullets_s4 = [
    ("Real-Time Webcam Viewfinder", "Kamera aktif secara instan dengan face guide oval untuk memastikan wajah teknisi tegak dan jelas."),
    ("Permanent Digital Watermark", "Foto otomatis dicap stempel digital permanen: NIK, Nama, Jam Detik WIB, Koordinat GPS, SSID WiFi, dan IP Subnet."),
    ("Auto Cut-Off Toleransi Shift", "Batas masuk 08:30 WIB. Keterlambatan terhitung otomatis per menit dan mewajibkan form alasan keterlambatan."),
    ("Geofencing Radius Kantor/Site", "Pengecekan koordinat fisik terhadap titik kantor pusat (Manyar JIIPE / Surabaya HQ) dengan toleransi 100m."),
    ("Audit Trail Log Presensi", "Setiap data presensi tersimpan permanen di database lokal/cloud lengkap dengan foto bukti selfie yang dapat diaudit.")
]

for title, desc in bullets_s4:
    p_b = tf_l.add_paragraph()
    p_b.space_before = Pt(8)
    r1 = p_b.add_run()
    r1.text = "• " + title + ": "
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(10.5)
    r2 = p_b.add_run()
    r2.text = desc
    r2.font.color.rgb = SLATE_LIGHT
    r2.font.size = Pt(10)

add_card(s4, 7.8, 1.7, 4.733, 4.3, border_color=ORANGE, bg_color=CARD_BG_ALT)
tb_r = s4.shapes.add_textbox(Inches(8.05), Inches(1.9), Inches(4.2), Inches(3.9))
tf_r = tb_r.text_frame
tf_r.word_wrap = True

p_r = tf_r.paragraphs[0]
p_r.text = "ANATOMI WATERMARK DIGITAL"
p_r.font.name = FONT_CODE
p_r.font.size = Pt(12)
p_r.font.bold = True
p_r.font.color.rgb = MINT

watermark_sample = (
    "---------------------------------------\n"
    "PRIME HRIS • VERIFIED ATTENDANCE STAMP\n"
    "---------------------------------------\n"
    "NIK      : DMJ-2026-003\n"
    "Nama     : Hendra Kurniawan\n"
    "Divisi   : Engineering & Fabrikasi\n"
    "Waktu    : 2026-09-09 08:04:12 WIB\n"
    "GPS      : -7.114210, 112.584102\n"
    "Lokasi   : Workshop Manyar JIIPE Gresik\n"
    "Network  : SSID: DMJ-Corporate-5G (OK)\n"
    "IP Pub   : 103.31.205.218 (Biznet Gio)\n"
    "Security : SHA256-TOKEN: a8f9e0c14b2d\n"
    "---------------------------------------\n"
    "STATUS   : TERVERIFIKASI TEPAT WAKTU"
)

p_w = tf_r.add_paragraph()
p_w.space_before = Pt(8)
p_w.text = watermark_sample
p_w.font.name = FONT_CODE
p_w.font.size = Pt(9)
p_w.font.color.rgb = SLATE_LIGHT

add_card(s4, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_n = s4.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_n = tb_n.text_frame
p_n = tf_n.paragraphs[0]
p_n.text = "[ ANTI-FRAUD GUARANTEE ] : Watermark dicetak langsung pada canvas bitmap gambar, mustahil dipalsukan dengan screenshot atau upload galeri biasa."
p_n.font.name = FONT_CODE
p_n.font.size = Pt(10)
p_n.font.bold = True
p_n.font.color.rgb = ORANGE_LIGHT

# ==============================================================================
# SLIDE 5: SEQUENTIAL WIFI GATEWAY & DINAS LUAR MENDADAK
# ==============================================================================
s5 = prs.slides.add_slide(blank_slide_layout)
apply_background(s5)
add_header(s5, "Modul 2 • Network Gate", "Sequential Corporate WiFi Gate & Opsi Dinas Luar Mendadak",
           "Perlindungan berlapis terhadap jaringan fiktif dengan fleksibilitas tinggi saat penugasan mendadak")
add_footer(s5, 5)

add_card(s5, 0.8, 1.7, 5.75, 4.3, border_color=BORDER_DARK)
tb_w = s5.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(5.25), Inches(3.9))
tf_w = tb_w.text_frame
tf_w.word_wrap = True

p = tf_w.paragraphs[0]
p.text = "SEQUENTIAL WIFI GATEWAY"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ORANGE

wifi_pts = [
    ("Tombol Submit Terkunci", "Tombol presensi WFO otomatis disabled jika perangkat belum terhubung ke Access Point kantor."),
    ("Validasi SSID & BSSID", "Mencocokkan nama SSID resmi ('DMJ-Corporate-5G') dan BSSID hardware router PT Dwi Martha Jaya."),
    ("IP Whitelist Subnet", "Memastikan IP Address berasal dari subnet ISP korporat terdaftar (misal Biznet Gio)."),
    ("Hardware MAC Binding", "Pengecekan MAC address perangkat resmi terdaftar untuk menangkal penyusup luar."),
    ("Anti-VPN & Fake GPS Shield", "Mendeteksi penggunaan proxy fiktif dan langsung menolak proses absensi.")
]

for t, d in wifi_pts:
    p_b = tf_w.add_paragraph()
    p_b.space_before = Pt(8)
    r1 = p_b.add_run()
    r1.text = "✔ " + t + ": "
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(10)
    r2 = p_b.add_run()
    r2.text = d
    r2.font.color.rgb = SLATE_LIGHT
    r2.font.size = Pt(9.5)

add_card(s5, 6.75, 1.7, 5.75, 4.3, border_color=MINT, bg_color=CARD_BG_ALT)
tb_d = s5.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.25), Inches(3.9))
tf_d = tb_d.text_frame
tf_d.word_wrap = True

p_d = tf_d.paragraphs[0]
p_d.text = "OPSI BYPASS : DINAS LUAR MENDADAK"
p_d.font.name = FONT_CODE
p_d.font.size = Pt(12)
p_d.font.bold = True
p_d.font.color.rgb = MINT

dinas_pts = [
    ("Kebutuhan Riil Lapangan", "Teknisi dipanggil darurat ke site klien (contoh: Manyar Smelter / Sorowako) tanpa sempat mampir kantor."),
    ("Bypass Terkendali (Controlled)", "Memungkinkan presensi dari jaringan seluler luar tanpa terikat WiFi kantor DMJ."),
    ("Wajib Input Data Penugasan", "Karyawan wajib memilih Klien Tujuan, Nomor SPK Proyek, Lokasi Site, dan Alasan Urgensi."),
    ("Routing Otomatis ke Approval", "Presensi dinas luar langsung masuk antrean 'Pending Approval' di portal Manager/Direksi."),
    ("Pencegahan Alpa Salah Catat", "Menghindarkan staf teladan dari hukuman alpa/potongan gaji akibat tugas mendadak.")
]

for t, d in dinas_pts:
    p_b = tf_d.add_paragraph()
    p_b.space_before = Pt(8)
    r1 = p_b.add_run()
    r1.text = "★ " + t + ": "
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(10)
    r2 = p_b.add_run()
    r2.text = d
    r2.font.color.rgb = SLATE_LIGHT
    r2.font.size = Pt(9.5)

add_card(s5, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_b = s5.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ KESEIMBANGAN SISTEM ] : Disiplin ketat bagi staf reguler kantor, sekaligus fleksibilitas tinggi bagi tim penugasan proyek lapangan."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = ORANGE_LIGHT

# ==============================================================================
# SLIDE 6: MASTER DATA SDM & MULTI-TIER APPROVAL HUB
# ==============================================================================
s6 = prs.slides.add_slide(blank_slide_layout)
apply_background(s6)
add_header(s6, "Modul 3 • SDM & Approval", "Master Data Karyawan & Pusat Persetujuan Terpadu",
           "Sentralisasi identitas SDM, struktur gaji, dan penyelesaian disposisi cuti/lembur satu pintu")
add_footer(s6, 6)

add_card(s6, 0.8, 1.7, 5.75, 4.3, border_color=BORDER_DARK)
tb_m = s6.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(5.25), Inches(3.9))
tf_m = tb_m.text_frame
tf_m.word_wrap = True

p = tf_m.paragraphs[0]
p.text = "DIREKTORI MASTER DATA KARYAWAN"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ORANGE

sdm_pts = [
    ("Klasifikasi Staf Lengkap", "Karyawan Tetap, Kontrak, PKWT Proyek, Staf Kantor Pusat, dan Teknisi Fabrikasi Workshop."),
    ("Parameter Kompensasi Gaji", "Gaji Pokok, Tunjangan Jabatan, Tunjangan Harian Makan & Transport, Tunjangan Site Remote."),
    ("Akun Bank Payroll Resmi", "Pencatatan nomor rekening BCA / Mandiri siap ekspor ke format bank transfer otomatis."),
    ("Tracking Saldo Cuti Tahunan", "Monitoring kuota cuti tahunan (12 hari), cuti bersama, dan riwayat pemakaian cuti."),
    ("Device Hardware Binding", "Perekaman MAC address perangkat seluler dan laptop resmi untuk keamanan otorisasi.")
]

for t, d in sdm_pts:
    p_b = tf_m.add_paragraph()
    p_b.space_before = Pt(8)
    r1 = p_b.add_run()
    r1.text = "▶ " + t + ": "
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(10)
    r2 = p_b.add_run()
    r2.text = d
    r2.font.color.rgb = SLATE_LIGHT
    r2.font.size = Pt(9.5)

add_card(s6, 6.75, 1.7, 5.75, 4.3, border_color=BORDER_DARK)
tb_a = s6.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.25), Inches(3.9))
tf_a = tb_a.text_frame
tf_a.word_wrap = True

p = tf_a.paragraphs[0]
p.text = "MULTI-TIER APPROVAL HUB"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = MINT

app_pts = [
    ("Disposisi Cuti & Izin Sakit", "Persetujuan atasan langsung yang otomatis memotong kuota cuti atau memverifikasi surat dokter."),
    ("Dispensasi Keterlambatan", "Verifikasi alasan keterlambatan (>08:30) untuk membebaskan penalti potongan gaji staf."),
    ("Validasi Penugasan Dinas Luar", "Pengecekan bukti foto selfie geotag dan SPK penugasan site darurat."),
    ("Klaim Reimbursement Proyek", "Persetujuan biaya bensin, tol, material darurat, dan uang makan site proyek."),
    ("Multi-Level Workflow", "Hierarki berjenjang: Pengajuan Staf -> Review Supervisor -> Approval Manajer/HR.")
]

for t, d in app_pts:
    p_b = tf_a.add_paragraph()
    p_b.space_before = Pt(8)
    r1 = p_b.add_run()
    r1.text = "✔ " + t + ": "
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(10)
    r2 = p_b.add_run()
    r2.text = d
    r2.font.color.rgb = SLATE_LIGHT
    r2.font.size = Pt(9.5)

add_card(s6, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_b = s6.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ AUDIT TRAIL ] : Setiap aksi approval tersimpan dengan stempel waktu detik, nama reviewer, catatan disposisi, dan status final."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = ORANGE_LIGHT

# ==============================================================================
# SLIDE 7: MESIN PENGGAJIAN & FORMULA LEMBUR DEPNAKER RI
# ==============================================================================
s7 = prs.slides.add_slide(blank_slide_layout)
apply_background(s7)
add_header(s7, "Modul 4 • Payroll Engine", "Otomatisasi Payroll & Perhitungan Lembur Standar Depnaker",
           "Perhitungan gaji akurat, potongan BPJS/PPh21, dan slip gaji resmi ber-QR Code validasi")
add_footer(s7, 7)

add_card(s7, 0.8, 1.7, 5.75, 4.3, border_color=ORANGE, bg_color=CARD_BG_ALT)
tb_p = s7.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(5.25), Inches(3.9))
tf_p = tb_p.text_frame
tf_p.word_wrap = True

p = tf_p.paragraphs[0]
p.text = "FORMULA BAKU KEMENAKERTRANS RI"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ORANGE

dep_pts = [
    ("Upah Lembur per Jam", "Formula standar regulasi RI: Gaji Pokok / 173 per jam kerja."),
    ("Multiplier Jam Lembur", "Hari Kerja Biasa: 1.5x upah sejam untuk jam ke-1, 2.0x upah sejam untuk jam ke-2 dan seterusnya."),
    ("Multiplier Hari Libur / Weekend", "Hari Libur Resmi: 2.0x hingga 3.0x upah per jam sesuai durasi penugasan."),
    ("Tunjangan Proyek Fleksibel", "Tunjangan Harian Kehadiran Proyek, Uang Makan Site, dan Tunjangan Luar Kota."),
    ("Potongan Disiplin Terukur", "Pemotongan proporsional jika keterlambatan tidak mendapatkan dispensasi atasan.")
]

for t, d in dep_pts:
    p_b = tf_p.add_paragraph()
    p_b.space_before = Pt(8)
    r1 = p_b.add_run()
    r1.text = "• " + t + ": "
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(10)
    r2 = p_b.add_run()
    r2.text = d
    r2.font.color.rgb = SLATE_LIGHT
    r2.font.size = Pt(9.5)

add_card(s7, 6.75, 1.7, 5.75, 4.3, border_color=BORDER_DARK)
tb_s = s7.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.25), Inches(3.9))
tf_s = tb_s.text_frame
tf_s.word_wrap = True

p = tf_s.paragraphs[0]
p.text = "KOMPONEN PAJAK, BPJS & SLIP RESMI"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = MINT

comp_pts = [
    ("BPJS Ketenagakerjaan (3%)", "Potongan otomatis JHT 2% dan Jaminan Pensiun 1% dari upah bruto."),
    ("BPJS Kesehatan (1%)", "Potongan iuran kesehatan karyawan sesuai ketentuan upah maksimal."),
    ("Estimasi Pajak PPh 21", "Perhitungan tarif pajak progresif UU HPP (TER PPh 21) otomatis per bulan."),
    ("Slip Gaji Digital Resmi", "Layout cetak berstandar korporat berlogo PT Dwi Martha Jaya."),
    ("Validasi QR Code Digital", "QR Code verifikasi keaslian slip gaji untuk kebutuhan perbankan staf."),
    ("Ekspor Format Payroll Bank", "Generate file CSV/Excel daftar transfer bank payroll (BCA & Mandiri).")
]

for t, d in comp_pts:
    p_b = tf_s.add_paragraph()
    p_b.space_before = Pt(6)
    r1 = p_b.add_run()
    r1.text = "✔ " + t + ": "
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(10)
    r2 = p_b.add_run()
    r2.text = d
    r2.font.color.rgb = SLATE_LIGHT
    r2.font.size = Pt(9)

add_card(s7, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_b = s7.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ BEBAS HUMAN ERROR ] : Menghilangkan 100% kesalahan hitung manual lembur excel dan dispute ketidakpuasan teknisi."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = MINT

# ==============================================================================
# SLIDE 8: PROJECT COST ACCOUNTING & FORECASTING EAC
# ==============================================================================
s8 = prs.slides.add_slide(blank_slide_layout)
apply_background(s8)
add_header(s8, "Modul 5 • Cost Accounting", "Penggajian Berbasis Proyek & Peramalan Anggaran (EAC)",
           "Distribusi biaya riil upah langsung per kontrak proyek dan deteksi dini risiko overbudget")
add_footer(s8, 8)

projects_enhanced = [
    ("PROYEK 1 : SMELTER MANYAR JIIPE",
     "Pekerjaan: Piping, Mechanical & Electrical Commissioning",
     "• Alokasi Teknisi : 60% Jam Kerja Total\n"
     "• Realisasi Gaji  : Rp 24.500.000 (Terbayar)\n"
     "• Lembur Depnaker : Rp 8.420.000 (148 Jam)\n"
     "• Pagu Anggaran   : Rp 45.000.000\n"
     "• Sisa Anggaran   : Rp 12.080.000",
     "STATUS : SEHAT (73% Terpakai)",
     0.73, MINT),
    ("PROYEK 2 : WORKSHOP FABRIKASI",
     "Pekerjaan: Fabrikasi Chute, Flange & Mesin Bubut/CNC",
     "• Alokasi Teknisi : 25% Jam Kerja Total\n"
     "• Realisasi Gaji  : Rp 12.800.000 (Terbayar)\n"
     "• Lembur Depnaker : Rp 3.150.000 (62 Jam)\n"
     "• Pagu Anggaran   : Rp 20.000.000\n"
     "• Sisa Anggaran   : Rp 4.050.000",
     "STATUS : SEHAT (79% Terpakai)",
     0.79, MINT),
    ("PROYEK 3 : SITE SOROWAKO OVERHAUL",
     "Pekerjaan: Emergency Overhaul Bearing & SAG Mill Liner",
     "• Alokasi Teknisi : 15% Jam Kerja Total\n"
     "• Realisasi Gaji  : Rp 9.200.000 (Terbayar)\n"
     "• Lembur Depnaker : Rp 5.800.000 (112 Jam High Burn)\n"
     "• Pagu Anggaran   : Rp 14.000.000\n"
     "• Proyeksi EAC    : Rp 15.000.000 (OVERRUN)",
     "STATUS : PERINGATAN (107% EAC)",
     1.07, RED_ACCENT)
]

for i, (p_title, p_scope, p_details, p_stat, pct, col) in enumerate(projects_enhanced):
    x = 0.8 + i * 3.98
    card = add_card(s8, x, 1.7, 3.8, 4.3, border_color=col, bg_color=CARD_BG_ALT if col==RED_ACCENT else CARD_BG)
    
    tb = s8.shapes.add_textbox(Inches(x + 0.2), Inches(1.85), Inches(3.4), Inches(3.9))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = p_title
    p.font.name = FONT_CODE
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    
    p_sc = tf.add_paragraph()
    p_sc.space_before = Pt(2)
    p_sc.text = p_scope
    p_sc.font.name = FONT_BODY
    p_sc.font.size = Pt(8.5)
    p_sc.font.color.rgb = SLATE_MUTED
    
    p_dt = tf.add_paragraph()
    p_dt.space_before = Pt(6)
    p_dt.text = p_details
    p_dt.font.name = FONT_CODE
    p_dt.font.size = Pt(9)
    p_dt.font.color.rgb = SLATE_LIGHT
    
    # Progress Bar Container
    bar_bg = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.2), Inches(5.15), Inches(3.4), Inches(0.22))
    bar_bg.fill.solid()
    bar_bg.fill.fore_color.rgb = RGBColor(14, 20, 36)
    bar_bg.line.color.rgb = BORDER_DARK
    
    # Progress Bar Fill
    fill_w = min(3.4, 3.4 * (pct if pct <= 1.0 else 1.0))
    bar_fill = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.2), Inches(5.15), Inches(fill_w), Inches(0.22))
    bar_fill.fill.solid()
    bar_fill.fill.fore_color.rgb = col
    bar_fill.line.color.rgb = col
    
    p_st = tf.add_paragraph()
    p_st.space_before = Pt(36)
    p_st.text = p_stat
    p_st.font.name = FONT_CODE
    p_st.font.size = Pt(10)
    p_st.font.bold = True
    p_st.font.color.rgb = col

add_card(s8, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_b = s8.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ ESTIMATE AT COMPLETION (EAC) ] : Memproyeksikan total estimasi biaya gaji teknisi hingga kontrak rampung, mencegah kerugian tak terduga."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = ORANGE_LIGHT

# ==============================================================================
# SLIDE 9: PILIHAN PAKET INVESTASI & ADD-ON MODULAR
# ==============================================================================
s9 = prs.slides.add_slide(blank_slide_layout)
apply_background(s9)
add_header(s9, "Skema Investasi", "Pilihan Paket Penawaran & Layanan Running Apps VPS",
           "Struktur biaya transparan dan efisien disesuaikan dengan skala operasional PT Dwi Martha Jaya")
add_footer(s9, 9)

add_card(s9, 0.8, 1.7, 4.3, 4.3, border_color=BORDER_DARK)
tb_b = s9.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(3.9), Inches(3.9))
tf_b = tb_b.text_frame
tf_b.word_wrap = True

p = tf_b.paragraphs[0]
p.text = "PAKET BASIC PRIME HRIS"
p.font.name = FONT_CODE
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = SLATE_MUTED

p_pr = tf_b.add_paragraph()
p_pr.text = "Rp 2.500.000"
p_pr.font.name = FONT_TITLE
p_pr.font.size = Pt(22)
p_pr.font.bold = True
p_pr.font.color.rgb = WHITE

p_sub = tf_b.add_paragraph()
p_sub.text = "One-Time Setup & Lisensi"
p_sub.font.name = FONT_CODE
p_sub.font.size = Pt(9.5)
p_sub.font.color.rgb = ORANGE_LIGHT

feat_b = [
    "Hingga 50 Karyawan Aktif",
    "Presensi Kamera Selfie & Watermark",
    "Dual Dashboard HRIS & Kas",
    "Multi-Tier Approval Hub",
    "Mesin Penggajian Depnaker (Slip DMJ)",
    "Layanan Running Apps: Rp 25k-50k/bln",
    "1 GB Storage App + 10 GB Backup"
]
for f in feat_b:
    pf = tf_b.add_paragraph()
    pf.space_before = Pt(4)
    pf.text = "✔ " + f
    pf.font.name = FONT_BODY
    pf.font.size = Pt(9)
    pf.font.color.rgb = SLATE_LIGHT

add_card(s9, 5.3, 1.6, 4.5, 4.5, border_color=ORANGE, bg_color=CARD_BG_ALT)
tb_r = s9.shapes.add_textbox(Inches(5.5), Inches(1.75), Inches(4.1), Inches(4.1))
tf_r = tb_r.text_frame
tf_r.word_wrap = True

p = tf_r.paragraphs[0]
p.text = "★ PAKET REKOMENDASI (FULL SUITE)"
p.font.name = FONT_CODE
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ORANGE

p_pr = tf_r.add_paragraph()
p_pr.text = "Rp 4.500.000"
p_pr.font.name = FONT_TITLE
p_pr.font.size = Pt(24)
p_pr.font.bold = True
p_pr.font.color.rgb = WHITE

p_sub = tf_r.add_paragraph()
p_sub.text = "One-Time All-in Setup & Lisensi"
p_sub.font.name = FONT_CODE
p_sub.font.size = Pt(9.5)
p_sub.font.color.rgb = MINT

feat_r = [
    "UNLIMITED Karyawan & Multi-Site",
    "Presensi Kamera Watermark Digital",
    "Termasuk: Sequential WiFi Gate Anti-Fraud",
    "Termasuk: Project Costing & Forecast EAC",
    "Mesin Penggajian Lengkap Depnaker",
    "Dual Dashboard + Multi-Tier Approval Hub",
    "Layanan Running Apps: Rp 25k-50k/bln",
    "Garansi Uptime 99.9% + Prioritas 12 Bulan"
]
for f in feat_r:
    pf = tf_r.add_paragraph()
    pf.space_before = Pt(4)
    pf.text = "★ " + f
    pf.font.name = FONT_BODY
    pf.font.size = Pt(9)
    pf.font.bold = True if "Termasuk" in f else False
    pf.font.color.rgb = WHITE if "Termasuk" in f else SLATE_LIGHT

add_card(s9, 10.0, 1.7, 2.533, 4.3, border_color=BORDER_DARK)
tb_a = s9.shapes.add_textbox(Inches(10.15), Inches(1.85), Inches(2.2), Inches(3.9))
tf_a = tb_a.text_frame
tf_a.word_wrap = True

p = tf_a.paragraphs[0]
p.text = "ADD-ON MODULAR"
p.font.name = FONT_CODE
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = MINT

addons = [
    ("Sequential WiFi Gate", "+ Rp 1.000.000"),
    ("Project Cost Accounting", "+ Rp 1.000.000"),
    ("PWA Mobile Android", "+ Rp 750.000"),
    ("Full Source Code", "+ Rp 2.500.000")
]
for name, prc in addons:
    pa = tf_a.add_paragraph()
    pa.space_before = Pt(8)
    r1 = pa.add_run()
    r1.text = name + "\n"
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.size = Pt(9)
    r2 = pa.add_run()
    r2.text = prc
    r2.font.color.rgb = ORANGE_LIGHT
    r2.font.size = Pt(8.5)

add_card(s9, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_b = s9.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ RUNNING APPS COMMITMENT ] : Biaya sewa server hanya Rp 25k – 50k / bulan, sudah mencakup domain HTTPS dan daily cloud backup."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = MINT

# ==============================================================================
# SLIDE 10: TIMELINE ROADMAP 5 MINGGU (AGILE)
# ==============================================================================
s10 = prs.slides.add_slide(blank_slide_layout)
apply_background(s10)
add_header(s10, "Timeline Implementasi", "Jadwal Pelaksanaan Proyek 5 Minggu (Agile Sprints)",
           "Metodologi implementasi bertahap dan cepat tanpa mengganggu aktivitas operasional berjalan")
add_footer(s10, 10)

sprints = [
    ("MINGGU 1", "Sprint 1: Inception", "• Kick-off meeting & requirement finalize\n• Whitelist SSID WiFi & BSSID kantor DMJ\n• Setup database karyawan, NIK & rekening\n• Konfigurasi domain & server VPS running"),
    ("MINGGU 2-3", "Sprint 2: Core Development", "• Implementasi kamera viewfinder & watermark\n• Penguncian Sequential WiFi Gateway\n• Integrasi alur bypass Dinas Luar Mendadak\n• Pembangunan mesin payroll lembur Depnaker"),
    ("MINGGU 4", "Sprint 3: Testing & UAT", "• User Acceptance Testing (UAT) staf & HR\n• Kalibrasi geofence Manyar JIIPE & Surabaya\n• Pengujian simulasi pembobolan Fake GPS/VPN\n• Verifikasi akurasi potongan BPJS & PPh21"),
    ("MINGGU 5", "Sprint 4: Go-Live Resmi", "• Pelatihan HR Admin & jajaran pimpinan\n• Handover dokumentasi & credential akses\n• Migrasi data absensi aktif PT DMJ\n• Go-Live resmi & aktivasi garansi SLA 12 bln")
]

for i, (m_tag, m_title, m_desc) in enumerate(sprints):
    x = 0.8 + i * 2.98
    card = add_card(s10, x, 1.8, 2.8, 4.1, border_color=ORANGE if i==3 else BORDER_DARK)
    
    tb = s10.shapes.add_textbox(Inches(x + 0.15), Inches(1.95), Inches(2.5), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = m_tag
    p.font.name = FONT_CODE
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    p2.text = m_title
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    
    p3 = tf.add_paragraph()
    p3.space_before = Pt(10)
    p3.text = m_desc
    p3.font.name = FONT_BODY
    p3.font.size = Pt(9.5)
    p3.font.color.rgb = SLATE_LIGHT

add_card(s10, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_b = s10.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ SPEED TO VALUE ] : Sistem siap beroperasi penuh dalam 35 hari kalender dengan jaminan pendampingan langsung oleh lead developer."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = MINT

# ==============================================================================
# SLIDE 11: SLA, KEAMANAN & KOMITMEN KUALITAS
# ==============================================================================
s11 = prs.slides.add_slide(blank_slide_layout)
apply_background(s11)
add_header(s11, "Jaminan Kualitas", "Service Level Agreement (SLA) & Standar Keamanan Data",
           "Kepastian keandalan sistem beroperasi 24/7 dengan perlindungan data kepegawaian standar industri")
add_footer(s11, 11)

sla_items = [
    ("99.9% Uptime Availability", 
     "Infrastruktur Cloud High-Availability",
     "Aplikasi berjalan di server Linux cloud dengan proteksi Cloudflare SSL dan reverse proxy ganda. Arsitektur stateless memastikan ketersediaan 24/7 tanpa downtime mengganggu jam absensi kritis."),
    ("Respon Cepat Gangguan",
     "SLA Respon < 2 Jam Kasus Kritis",
     "Kategori Kritis (Aplikasi down/presensi gagal total): Respon teknis < 2 Jam.\nKategori Moderat (Revisi fitur/data): Respon < 6 Jam.\nDidukung jalur komunikasi langsung via hotline WhatsApp pengembang."),
    ("Backup & Proteksi Data",
     "Automated Daily Incremental Backup",
     "Seluruh data kehadiran, foto watermark, dan master gaji dicadangkan harian secara otomatis ke cloud storage terpisah. Enkripsi data at-rest dan in-transit menggunakan sertifikat SSL 256-bit.")
]

for i, (tag, title, desc) in enumerate(sla_items):
    x = 0.8 + i * 3.98
    card = add_card(s11, x, 1.8, 3.8, 4.1, border_color=BORDER_DARK)
    
    tb = s11.shapes.add_textbox(Inches(x + 0.2), Inches(1.95), Inches(3.4), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = tag
    p.font.name = FONT_CODE
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = MINT
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    p2.text = title
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    
    p3 = tf.add_paragraph()
    p3.space_before = Pt(10)
    p3.text = desc
    p3.font.name = FONT_BODY
    p3.font.size = Pt(10)
    p3.font.color.rgb = SLATE_LIGHT

add_card(s11, 0.8, 6.2, 11.733, 0.65, border_color=ORANGE, bg_color=CARD_BG_ALT)
tb_b = s11.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ 12 BULAN GARANSI BEBAS BUG ] : Pemeliharaan penuh terhadap perbaikan bug dan penyesuaian regulasi ketenagakerjaan selama masa kontrak."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = WHITE

# ==============================================================================
# SLIDE 12: PENUTUP, KONTAK & LEMBAR PENGESAHAN
# ==============================================================================
s12 = prs.slides.add_slide(blank_slide_layout)
apply_background(s12)
add_header(s12, "Penutup & Pengesahan", "Akses Live Demo & Lembar Pengesahan Kerjasama",
           "Langkah awal modernisasi tata kelola SDM industri PT Dwi Martha Jaya")
add_footer(s12, 12)

add_card(s12, 0.8, 1.7, 5.75, 4.3, border_color=ORANGE, bg_color=CARD_BG_ALT)
tb_c = s12.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(5.25), Inches(3.9))
tf_c = tb_c.text_frame
tf_c.word_wrap = True

p = tf_c.paragraphs[0]
p.text = "AKSES PROTOTIPE & KONTAK DIRECT"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ORANGE

contact_info = (
    "• Platform Showcase Live : https://dmj.primeprojectx.net\n"
    "  (Dilengkapi Interactive Tour Demo & Typewriter Walkthrough)\n\n"
    "• Vendor Pengembang      : PT PRIME INFINITY SYSTEMS\n"
    "• Lead Platform Architect : Galih Primananda, S.E.\n"
    "• WhatsApp Langsung      : +62 895 2425 7778\n"
    "• Email Resmi            : galih@primeprojectx.net\n"
    "• Platform Ekosistem     : https://primeprojectx.net\n"
    "• Domisili Operasional   : Surabaya & Jakarta, Indonesia\n\n"
    "Silakan buka web apps atau hubungi WhatsApp kami untuk diskusi kick-off."
)

p_ci = tf_c.add_paragraph()
p_ci.space_before = Pt(6)
p_ci.text = contact_info
p_ci.font.name = FONT_CODE
p_ci.font.size = Pt(9.5)
p_ci.font.color.rgb = WHITE

add_card(s12, 6.75, 1.7, 5.75, 4.3, border_color=BORDER_DARK)
tb_s = s12.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.25), Inches(3.9))
tf_s = tb_s.text_frame
tf_s.word_wrap = True

p = tf_s.paragraphs[0]
p.text = "LEMBAR PENGESAHAN DOKUMEN"
p.font.name = FONT_CODE
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = MINT

sign_text = (
    "Diajukan Oleh:\n"
    "PT PRIME INFINITY SYSTEMS (PRIME PROJECTX)\n\n"
    "(Tanda Tangan & Cap Digital)\n\n"
    "Galih Primananda, S.E.\n"
    "Managing Director & Lead Platform Architect\n"
    "----------------------------------------------------\n"
    "Disetujui & Diterima Oleh:\n"
    "PT DWI MARTHA JAYA\n\n"
    "(Tanda Tangan & Cap Perusahaan)\n\n"
    "( ................................................................ )\n"
    "Authorized Director / General Manager"
)

p_si = tf_s.add_paragraph()
p_si.space_before = Pt(6)
p_si.text = sign_text
p_si.font.name = FONT_CODE
p_si.font.size = Pt(9.5)
p_si.font.color.rgb = SLATE_LIGHT

add_card(s12, 0.8, 6.2, 11.733, 0.65, border_color=BORDER_DARK)
tb_b = s12.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
tf_b = tb_b.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "[ KOMITMEN KAMI ] : Menghadirkan solusi teknologi mutakhir yang tepat guna, efisien, dan memberikan nilai tambah nyata bagi PT Dwi Martha Jaya."
p_b.font.name = FONT_CODE
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = MINT

# Save presentation
output_pptx = os.path.join(r"C:\Users\PRIMA\timesweet-hris", "PROPOSAL_PRIME_HRIS_PT_DWI_MARTHA_JAYA.pptx")
prs.save(output_pptx)
print("SUCCESS: Saved presentation to", output_pptx)
